import os
import time
import requests
import pandas as pd
from docx import Document

BASE_URL = "http://127.0.0.1:8000"

def create_dummy_files():
    """Stvara privremene testne datoteke raznih formata."""
    os.makedirs("./test_docs", exist_ok=True)
    
    # 1. TXT
    with open("./test_docs/test_info.txt", "w", encoding="utf-8") as f:
        f.write("Sustav za upravljanje dronovima koristi protokol MAVLink za komunikaciju.")

    # 2. Markdown
    with open("./test_docs/test_notes.md", "w", encoding="utf-8") as f:
        f.write("# Projekt Dronovi\nKlasifikacija slika iz zraka vrši się pomoću YOLOv8 modela.")

    # 3. DOCX (Word)
    doc = Document()
    doc.add_heading("Specifikacija Projekta", level=1)
    doc.add_paragraph("Anotacija podataka za AI modele izvodi se u alatu Label Studio.")
    doc.save("./test_docs/test_doc.docx")

    # 4. XLSX (Excel)
    df = pd.DataFrame({
        "Modul": ["Anotacija", "Dronovi", "RAG Pipeline"],
        "Voditelj": ["Tomislav", "Ana", "Marko"],
        "Budzet_EUR": [15000, 25000, 10000]
    })
    df.to_excel("./test_docs/test_budget.xlsx", index=False)

def test_healthcheck():
    print("\n[TEST 1] Provjera dostupnosti API-ja...")
    response = requests.get(f"{BASE_URL}/")
    assert response.status_code == 200
    print(" Status 200 OK -> API je aktivan!")

def test_upload_endpoints():
    print("\n[TEST 2] Testiranje slanja raznih formata (/api/v1/upload)...")
    test_files = [
        "./test_docs/test_info.txt",
        "./test_docs/test_notes.md",
        "./test_docs/test_doc.docx",
        "./test_docs/test_budget.xlsx"
    ]

    for file_path in test_files:
        filename = os.path.basename(file_path)
        with open(file_path, "rb") as f:
            files = {"file": (filename, f)}
            response = requests.post(f"{BASE_URL}/api/v1/upload", files=files)
            
            assert response.status_code == 202
            print(f" Datoteka {filename} uspesno poslata (Status 202 Accepted).")

def test_reindex_endpoint():
    print("\n[TEST 3] Testiranje reindeksiranja (/api/v1/reindex)...")
    response = requests.post(f"{BASE_URL}/api/v1/reindex")
    assert response.status_code == 202
    print(" Reindeksiranje pokrenuto u pozadini (Status 202 Accepted).")

def test_query_rag():
    print("\n[TEST 4] Čekanje 5s da pozadinski zadaci završe i slanje upita (/api/query)...")
    time.sleep(5)  # Dajemo pozadinskom radniku malo vremena da vektorizira datoteke

    queries = [
        "Koji se alat koristi za anotaciju podataka?",
        "Koliki je budžet za modul Dronovi?",
        "Koji se protokol koristi za komunikaciju s dronovima?"
    ]

    for q in queries:
        response = requests.post(f"{BASE_URL}/api/query", json={"question": q})
        assert response.status_code == 200
        data = response.json()
        print(f"\nPitanje: {q}")
        print(f"Odgovor: {data['answer']}")
        print(f"Izvori:  {data['sources']}")

if __name__ == "__main__":
    print("=== Pokretanje E2E API Testova ===")
    try:
        create_dummy_files()
        test_healthcheck()
        test_upload_endpoints()
        test_reindex_endpoint()
        test_query_rag()
        print("\n Svi testovi su uspješno prošli!")
    except requests.exceptions.ConnectionError:
        print("\n GREŠKA: FastAPI poslužitelj nije pokrenut! Pokrenite 'uvicorn main:app --reload' u drugom terminalu.")
    except AssertionError as e:
        print(f"\n TEST NIJE PROŠAO: Neočekivan status odziva sa poslužitelja.")